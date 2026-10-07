"""Builds Conveyor_BOM_New_Build.xlsx from the original BOM plus the purchase invoices.

Run from anywhere:  /usr/bin/python3 make_bom.py
The original BOM in original_BOM/ is only read, never changed.
To add a new bill: add its lines to an invoice list below, map them in BOUGHT, and run again.
"""
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "original_BOM", "PLC_Control_Box_Revised_Final_BOM-1_MCB_Corrected.xlsx")
OUT = os.path.join(HERE, "Conveyor_BOM_New_Build.xlsx")

PAVITRA = "Pavitra Electricals, inv 7964/2026-27, 03-10-2026"
NANO = "Nanologic Automation, inv NA/SL/26-27/225, 05-10-2026"

# Pavitra Electricals and Industrial Suppliers, invoice 7964/2026-27, 03-Oct-2026 (2 pages).
# Read from the photo. Check: amounts add up to the printed basic total 14,145.00 and qty total 175.
PAVITRA_LINES = [
    (1, "3TIER 220V AC TOWER LIGHT WITH BUZZER", "85361010", 1, "Nos", 1350),
    (2, "P G 21 CABLE GLAND", "85369090", 5, "Nos", 30),
    (3, "P G 16 CABLE GLAND", "85369090", 5, "Nos", 20),
    (4, "P G 29 HOSE GLAND", "85369090", 2, "Nos", 75),
    (5, "P G 29 HOSE PIPE", "39172190", 2, "Mtr", 95),
    (6, "16 A 2P ON / OFF 61198SCB03TDYR", "85365090", 1, "Nos", 380),
    (7, "E M R SWITCH S/N", "85365020", 1, "Nos", 290),
    (8, "ILPB", "85381010", 1, "Nos", 450),
    (9, "PUSH BUTTON", "85369090", 2, "Pcs", 150),
    (10, "2 P SWITCH", "85381010", 1, "Nos", 25),
    (11, "6 A SPN MCB", "85362030", 1, "Nos", 350),
    (12, "10 A SPN MCB S/N", "85362030", 1, "Nos", 550),
    (13, "2 A SPN MCB", "85362030", 1, "Nos", 850),
    (14, "200 MM CABLE TIE (P)", "39232990", 1, "Pkt", 120),
    (15, "SMPS 5a 24 V Dc", "85045090", 1, "Nos", 1250),
    (16, "CX4U", "85389000", 10, "Nos", 25),
    (17, "CX 2.5 U", "85389000", 50, "Nos", 15),
    (18, "END PLATE Epcx2.5", "85389000", 10, "Nos", 8),
    (19, "END PLATE Epcx4", "85389000", 5, "Nos", 10),
    (20, "END CLAMP", "85389000", 10, "Nos", 10),
    (21, "1.5 SQ BUTTED LUGS (P)", "85369090", 2, "Pkt", 55),
    (22, "5 A 12 WAY CONNECTOR", "85366910", 1, "Nos", 65),
    (23, "15 A SS COMMAND BOX", "85361010", 1, "Nos", 150),
    (24, "15 A 3 PIN TOP", "85369090", 1, "Nos", 80),
    (25, "SHORT LINK", "85389000", 20, "Pcs", 95),
    (26, "0.5 SQ MM 5 CORE ORBIT R/F FRLS IND. CABLE 1.1KV", "85446020", 3, "Mtr", 70),
    (27, "0.5 SQ MM 12 CORE WIRE (M)", "85446020", 3, "Mtr", 130),
    (28, "1.5 SQ MM 3 CORE WIRE (M)", "85446020", 3, "Mtr", 90),
    (29, "0.5 SQ MM WIRE (C) 90mtr", "85446020", 1, "Coil", 995),
    (30, "1.5 SQ MM WIRE (M)", "85446020", 20, "Mtr", 35),
    (31, "DLINK: PATCHCARD : 1MTR CAT6 PATCHCORD GREY", "85446020", 1, "Nos", 150),
    (32, "DLINK: PATCHCARD : 2 MTR CAT6 PATCHCORD GREY", "85446020", 1, "Nos", 250),
    (33, "DLINK: PATCHCARD : 3 MTR CAT6 PATCHCORD GREY", "85446020", 1, "Nos", 380),
    (34, "MCB CHANNEL (M)", "85389000", 2, "Mtr", 60),
    (35, "B 40X40 PVC CABLE DUCT", "85389000", 4, "Mtr", 160),
]
assert sum(q * r for _, _, _, q, _, r in PAVITRA_LINES) == 14145
assert sum(q for _, _, _, q, _, _ in PAVITRA_LINES) == 175

# Nanologic Automation, invoice NA/SL/26-27/225, 05-Oct-2026. Two of each were bought;
# only ONE of each belongs to this machine (the second set is extra, not part of this BOM).
NANO_LINES = [
    (1, "6ES7214-1AG40-0XB0 SIMATIC S7-1200, CPU 1214C, compact CPU, DC/DC/DC", "85371000", 2, "PCS", 28000,
     "Serials SC-U6M113QV, SC-U6M1CXK4"),
    (2, "6AV2123-2GB03-0AX0 SIMATIC HMI, KTP700 Basic, 7\" TFT", "85371000", 2, "PCS", 30000,
     "Serials SC-T3JK1131, SC-T1GD2473"),
    (3, "FREIGHT CHARGES", "00996812", 1, "PCS", 450, ""),
]

# Old-BOM component -> (bought?, supplier, invoice item text, qty for THIS machine, rate ex GST, note)
OTHER = "Other shop - bill not given yet"
BOUGHT = {
    "PLC": ("Yes", NANO, "item 1: 6ES7214-1AG40-0XB0", 1, 28000,
            "2 bought, 1 for this machine. Write here which serial is fitted: SC-U6M113QV or SC-U6M1CXK4."),
    "24VDC power supply": ("Yes", PAVITRA, "item 15: SMPS 5a 24 V Dc", 1, 1250,
                           "Make not on bill (old = Multispan PS24-5A-P, 120 W). Check label: 24 V DC, at least 5 A."),
    "2P MCB C10": ("Yes", PAVITRA, "item 12: 10 A SPN MCB S/N", 1, 550,
                   "Bill says SPN (single pole + neutral); old is Schneider 2-pole C10. Check it is C-curve."),
    "2P MCB C6": ("Yes", PAVITRA, "item 11: 6 A SPN MCB", 1, 350,
                  "Bill says SPN; old is Schneider 2-pole C6. Check it is C-curve."),
    "2P MCB C2": ("Yes", PAVITRA, "item 13: 2 A SPN MCB", 1, 850,
                  "Bill says SPN; old is Schneider 2-pole C2. Check it is C-curve."),
    "VFD": ("?", OTHER, "", 1, None, "Delta VFD004EL21W-1. Give me the bill."),
    "HMI": ("Yes", NANO, "item 2: 6AV2123-2GB03-0AX0", 1, 30000,
            "2 bought, 1 for this machine. Write here which serial is fitted: SC-T3JK1131 or SC-T1GD2473."),
    "Rotary ON/OFF isolator": ("Yes", PAVITRA, "item 6: 16 A 2P ON / OFF 61198SCB03TDYR", 1, 380,
                               "Different make from old Salzer, same rating (16 A, 2 pole). OK."),
    "Industrial Ethernet switch": ("?", OTHER, "", 1, None, "Teltonika TSW010. Give me the bill."),
    "3-colour signal tower with buzzer": ("Yes - CHECK", PAVITRA, "item 1: 3TIER 220V AC TOWER LIGHT WITH BUZZER", 1, 1350,
                                          "WARNING: bill says 220 V AC. The PLC lamp outputs Q0.3-Q0.5 give 24 V DC. "
                                          "Check the voltage on the OLD tower light. If the old one is 24 V DC, exchange this one."),
    "Cycle START pushbutton": ("Yes", PAVITRA, "item 9: PUSH BUTTON (2 pcs, with STOP)", 1, 150,
                               "Colour not on bill. Old START is green."),
    "Cycle STOP pushbutton": ("Yes", PAVITRA, "item 9: PUSH BUTTON (2 pcs, with START)", 1, 150,
                              "Colour not on bill. Old STOP is orange."),
    "Manual/AUTO selector switch": ("Yes - CHECK", PAVITRA, "item 10: 2 P SWITCH", 1, 25,
                                    "Only Rs 25. Check it is a 22 mm 2-position panel selector, not a small toggle switch."),
    "VFD HEALTHY pilot lamp": ("Yes - CHECK", PAVITRA, "item 8: ILPB (illuminated push button)", 1, 450,
                               "Check colour (green) and lamp voltage: must be 24 V DC (fed from the VFD relay RB)."),
    "Emergency-stop pushbutton": ("Yes", PAVITRA, "item 7: E M R SWITCH S/N", 1, 290, ""),
    "Panel-mounted electrical socket": ("Yes", PAVITRA, "items 23 + 24: 15 A SS COMMAND BOX + 15 A 3 PIN TOP", 1, 230,
                                        "Socket box (Rs 150) + plug top (Rs 80)."),
    "Contact blocks": ("?", "", "", None, None, "Not on the bills. Check that the push buttons came with contact blocks."),
    "35mm slotted DIN rail": ("Yes - CHECK", PAVITRA, "item 34: MCB CHANNEL (M), 2 m", 2, 60,
                              "Bought 2 m; old BOM planned 4 m. Measure the back plate."),
    "Slotted wiring duct with cap ": ("Yes", PAVITRA, "item 35: B 40X40 PVC CABLE DUCT, 4 m", 4, 160,
                                      "40 x 40 mm. Old BOM planned 3 m."),
    "DIN-rail terminal blocks": ("Yes - CHECK", PAVITRA, "items 16 + 17: CX4U x10, CX 2.5 U x50", 60, None,
                                 "60 bought (old BOM planned 70). No green/yellow earth (PE) terminals on the bill."),
    "5-way screw terminal strip": ("Yes", PAVITRA, "item 22: 5 A 12 WAY CONNECTOR", 1, 65, "12-way strip; cut to size."),
    "White nylon cable glands": ("Yes", PAVITRA, "item 2: P G 21 CABLE GLAND", 5, 30, ""),
    "Small white cable gland": ("Yes", PAVITRA, "item 3: P G 16 CABLE GLAND", 5, 20, ""),
    "DIN-rail end clamps/end stops": ("Yes", PAVITRA, "item 20: END CLAMP", 10, 10,
                                      "Also bought terminal end plates: items 18 and 19."),
    "Black cable gland": ("?", "", "", None, None, "Not on the bills."),
    "Flexible corrugated conduit": ("Yes", PAVITRA, "item 5: P G 29 HOSE PIPE, 2 m (+ item 4: 2 hose glands)", 2, 95, ""),
    "Internal mounting/back plate": ("?", OTHER, "", 1, None, "Not on the bills."),
    "Bootlace ferrules": ("Yes - CHECK", PAVITRA, "item 21: 1.5 SQ BUTTED LUGS (P), 2 packets", 2, 55,
                          "Lugs for 1.5 sq mm only. Ferrules for the 0.5 sq mm wires are not on the bill."),
    "Panel wire": ("Yes", PAVITRA, "items 29 + 30: 0.5 sq mm 90 m coil, 1.5 sq mm 20 m", None, None,
                   "Multi-core cables bought too: items 26, 27, 28."),
    "Wire/terminal markers": ("?", "", "", None, None, "Not on the bills."),
    "Cable ties": ("Yes", PAVITRA, "item 14: 200 MM CABLE TIE (P)", 1, 120, ""),
    "Panel mounting hardware": ("?", "", "", None, None, "Not on the bills."),
    "Conveyor motor": ("?", OTHER, "", 1, None, "Bonfiglioli BE 71B 4. Give me the bill."),
    "Photoelectric sensor": ("?", OTHER, "", 1, None, "Give me the bill and the model."),
    "Retro-reflector": ("?", OTHER, "", 1, None, "Give me the bill and the model."),
    "Logistics code reader / camera": ("?", OTHER, "", 1, None, "HIKROBOT MV-ID6200M-00C-NNG. Give me the bill."),
    "HIKROBOT logistics light": ("?", OTHER, "", 1, None, "Give me the bill and the label."),
    "C-mount 12mm machine-vision lens": ("?", OTHER, "", 1, None, "Give me the bill and the lens marking."),
    "Vision-system SMPS / power supply": ("?", OTHER, "", None, None, "Only if the old machine has a separate camera supply."),
}

# Bought items that are not in the old BOM.
EXTRA = [
    ("Wiring", "Terminal end plates", "Connectwell (EPCX)", "EPCX2.5 x10, EPCX4 x5", "End covers for CX terminal blocks", "15",
     PAVITRA, "items 18 + 19", None),
    ("Wiring", "Terminal short links (jumpers)", "", "SHORT LINK", "Bridges between terminal blocks", "20",
     PAVITRA, "item 25", 95),
    ("Wiring", "0.5 sq mm 5-core cable", "Orbit", "R/F FRLS 1.1 kV", "Multi-core cable", "3 m", PAVITRA, "item 26", 70),
    ("Wiring", "0.5 sq mm 12-core cable", "", "", "Multi-core cable", "3 m", PAVITRA, "item 27", 130),
    ("Wiring", "1.5 sq mm 3-core cable", "", "", "Multi-core cable (230 V supply)", "3 m", PAVITRA, "item 28", 90),
    ("Networking", "Ethernet patch cord 1 m", "D-Link", "Cat6, grey", "PLC / HMI / camera / PC to switch", "1", PAVITRA, "item 31", 150),
    ("Networking", "Ethernet patch cord 2 m", "D-Link", "Cat6, grey", "PLC / HMI / camera / PC to switch", "1", PAVITRA, "item 32", 250),
    ("Networking", "Ethernet patch cord 3 m", "D-Link", "Cat6, grey", "PLC / HMI / camera / PC to switch", "1", PAVITRA, "item 33", 380),
    ("Tool (not in panel)", "USB to RS-485 converter", "Sharvi", "ISO-B-RS485-A (FTDI FT232RL, isolated)",
     "Used to read the VFD parameters from the laptop", "1", OTHER, "", None),
]

BLUE = PatternFill("solid", fgColor="FF5B9BD5")
GREEN = PatternFill("solid", fgColor="FFE2F0D9")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
ORANGE = PatternFill("solid", fgColor="FFFFC000")
GREY = PatternFill("solid", fgColor="FFEDEDED")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row, i, n)
        c.font, c.fill, c.alignment = BOLD, BLUE, WRAP
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


wb = openpyxl.load_workbook(SRC)
bom = wb.active
bom.title = "BOM new build"
bom["A1"] = "Conveyor panel BOM - new build. Columns A-H = original BOM (unchanged). Columns I-N = what was bought. " \
            "Extra PLC + HMI (second set) NOT included."
bom["A1"].font = BOLD
new_cols = ["Bought for new build?", "Supplier / invoice", "Item on invoice", "Qty for this machine",
            "Rate ex GST (Rs)", "Check / note"]
for i, (name, w) in enumerate(zip(new_cols, [16, 34, 40, 12, 12, 60]), 9):
    c = bom.cell(4, i, name)
    c.font, c.fill, c.alignment = BOLD, ORANGE, WRAP
    bom.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

seen = set()
last = 4
for r in range(5, bom.max_row + 1):
    comp = bom.cell(r, 2).value
    if not comp:
        continue
    last = r
    info = BOUGHT.get(comp)
    if info is None:
        raise SystemExit(f"BOM row {r} '{comp}' has no purchase mapping - add it to BOUGHT")
    seen.add(comp)
    for i, v in enumerate(info, 9):
        c = bom.cell(r, i, v)
        c.alignment = WRAP
        if "CHECK" in str(info[0]):
            c.fill = YELLOW
        elif info[0] == "Yes":
            c.fill = GREEN
        else:
            c.fill = GREY
missing = set(BOUGHT) - seen
assert not missing, missing

r = last + 2
bom.cell(r, 1, "Extra items bought that were not in the old BOM").font = BOLD
for sec, comp, mfr, model, spec, qty, sup, item, rate in EXTRA:
    r += 1
    vals = [sec, comp, mfr, model, spec, qty, "New", "", "Yes" if sup != OTHER else "Yes (bill not given)",
            sup, item, qty, rate, ""]
    for i, v in enumerate(vals, 1):
        c = bom.cell(r, i, v)
        c.alignment = WRAP
        if i >= 9:
            c.fill = GREEN
r += 2
bom.cell(r, 1, "Colours in columns I-N:").font = BOLD
for txt, fill in [("bought, matches", GREEN), ("bought, but CHECK the note", YELLOW),
                  ("not on the bills yet / other shop", GREY)]:
    r += 1
    c = bom.cell(r, 2, txt)
    c.fill = fill

# Invoice sheets, line by line as printed.
ws = wb.create_sheet("Invoice Pavitra 7964")
ws["A1"] = "Pavitra Electricals and Industrial Suppliers - Tax invoice 7964/2026-27 - 03-Oct-2026 - buyer Simpel Techlabs Pvt Ltd"
ws["A1"].font = BOLD
header(ws, 3, ["S.N", "Description (as printed)", "HSN", "Qty", "Unit", "Rate (Rs)", "Amount (Rs)"],
       [6, 52, 12, 8, 8, 12, 14])
for i, line in enumerate(PAVITRA_LINES, 4):
    sn, desc, hsn, q, u, rate = line
    for j, v in enumerate([sn, desc, hsn, q, u, rate, q * rate], 1):
        ws.cell(i, j, v)
n = 4 + len(PAVITRA_LINES)
for label, val in [("Basic total", 14145), ("CGST 9%", 1273.05), ("SGST 9%", 1273.05), ("Round off", -0.10),
                   ("Grand total (printed)", 16691.00)]:
    ws.cell(n, 2, label).font = BOLD
    ws.cell(n, 7, val).font = BOLD
    n += 1
ws.cell(n + 1, 2, "Check: line amounts add up to 14,145 and quantities to 175, same as printed on the bill.")

ws = wb.create_sheet("Invoice Nanologic 225")
ws["A1"] = "Nanologic Automation - Tax invoice NA/SL/26-27/225 - 05-Oct-2026 - PO 1420 - buyer Simpel Techlabs Pvt Ltd"
ws["A1"].font = BOLD
header(ws, 3, ["S.No", "Item (as printed)", "HSN", "Qty", "Unit", "Rate (Rs)", "Amount ex GST (Rs)", "Serial numbers"],
       [6, 60, 12, 6, 6, 12, 16, 36])
for i, (sn, desc, hsn, q, u, rate, ser) in enumerate(NANO_LINES, 4):
    for j, v in enumerate([sn, desc, hsn, q, u, rate, q * rate, ser], 1):
        ws.cell(i, j, v)
ws.cell(8, 2, "Taxable 1,16,450 + CGST 10,480.50 + SGST 10,480.50 = Total 1,37,411 (paid, balance 0)").font = BOLD
ws.cell(9, 2, "2 PLC + 2 HMI were bought. Only 1 PLC + 1 HMI belong to this machine; the second set is extra "
              "and is NOT counted in the BOM.")

wb.save(OUT)
print("saved", OUT)
